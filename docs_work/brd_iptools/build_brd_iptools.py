from __future__ import annotations

import os
import shutil
import tempfile
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from lxml import etree
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\WEB\laravel-inertia")
WORK = ROOT / "docs_work" / "brd_iptools"
REFERENCE = WORK / "BRD InfraSight reference.docx"
OUTPUT = ROOT / "outputs" / "BRD-IPTOOLS-LITE.docx"

BLACK = "000000"
DARK_BLUE = "1F4E79"
NAVY = "17365D"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F2F7FB"
LIGHT_GRAY = "D9D9D9"
MID_GRAY = "696969"
WHITE = "FFFFFF"


def clear_element(element):
    for child in list(element):
        element.remove(child)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_font(run, name="Arial", size=9, bold=None, italic=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.04

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(5)

    for name, size, before, after in [
        ("Heading 1", 14, 5, 4),
        ("Heading 2", 11, 4, 2),
    ]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("Header", "Footer", "List Paragraph"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")


def add_body(doc, text="", bold=False, italic=False, color=BLACK, size=9, align=None, after=3):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.04
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_note(doc, text):
    return add_body(doc, text, italic=True, color=MID_GRAY, size=8.5, after=4)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.45)
    p.paragraph_format.first_line_indent = Cm(-0.22)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_font(r, size=9)
    return p


def add_number(doc, text, number):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.3)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(f"{number}. {text}")
    set_font(r, size=8.5)
    return p


def add_heading(doc, text, level=1, new_page=False):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    if new_page:
        p.paragraph_format.page_break_before = True
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.1, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    set_cant_split(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(str(value))
        set_font(r, size=font_size, bold=True, color=BLACK)
        if widths:
            cell.width = Cm(widths[idx])
    for r_idx, row_data in enumerate(rows):
        row = table.add_row()
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        if len(rows) <= 12:
            set_cant_split(row)
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            if r_idx % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if idx == 0 and len(headers) > 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run("" if value is None else str(value))
            set_font(r, size=font_size, color=BLACK)
            if widths:
                cell.width = Cm(widths[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_field(paragraph, instruction: str, cached_text=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = cached_text
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    return run


def add_toc(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    add_field(p, 'TOC \\o "1-2" \\h \\z \\u', "Daftar isi akan diperbarui otomatis.")


def font_path(bold=False):
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    return None


def draw_centered_text(draw, box, text, font, fill):
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=4, align="center")
    w = bbox[2] - bbox[0]
    h = bbox[3] - bbox[1]
    draw.multiline_text(((x1 + x2 - w) / 2, (y1 + y2 - h) / 2), text, font=font, fill=fill, spacing=4, align="center")


def arrow(draw, start, end, color=NAVY, width=5):
    draw.line([start, end], fill=f"#{color}", width=width)
    x2, y2 = end
    x1, y1 = start
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    size = 16
    p1 = (x2, y2)
    p2 = (x2 - ux * size + px * size * 0.55, y2 - uy * size + py * size * 0.55)
    p3 = (x2 - ux * size - px * size * 0.55, y2 - uy * size - py * size * 0.55)
    draw.polygon([p1, p2, p3], fill=f"#{color}")


def build_diagrams():
    bold = ImageFont.truetype(font_path(True), 28)
    regular = ImageFont.truetype(font_path(False), 23)
    small = ImageFont.truetype(font_path(False), 19)

    current = Image.new("RGB", (1600, 720), "white")
    d = ImageDraw.Draw(current)
    d.text((65, 35), "KONDISI PROSES PEMBANDING SEBELUM APLIKASI", font=bold, fill=f"#{NAVY}")
    labels = [
        ("1", "Permintaan\nlewat pesan\natau formulir"),
        ("2", "Cek stok\npada catatan\nterpisah"),
        ("3", "Persetujuan\ntidak berada\ndi satu alur"),
        ("4", "Serah terima\ndan foto\nterpisah"),
        ("5", "Rekap mutasi\ndan laporan\nmanual"),
    ]
    start_x, gap, box_w, box_h, top = 70, 45, 250, 250, 170
    for i, (num, label) in enumerate(labels):
        x = start_x + i * (box_w + gap)
        d.rounded_rectangle((x, top, x + box_w, top + box_h), radius=22, fill="#F2F7FB", outline=f"#{DARK_BLUE}", width=4)
        d.ellipse((x + 92, top + 18, x + 158, top + 84), fill=f"#{DARK_BLUE}")
        draw_centered_text(d, (x + 92, top + 18, x + 158, top + 84), num, bold, "white")
        draw_centered_text(d, (x + 16, top + 92, x + box_w - 16, top + box_h - 18), label, regular, "#111111")
        if i < len(labels) - 1:
            arrow(d, (x + box_w + 8, top + box_h / 2), (x + box_w + gap - 8, top + box_h / 2))
    d.text((70, 505), "Risiko utama: stok tidak seragam, approval sulit ditelusuri, bukti tersebar, dan rekap terlambat.", font=regular, fill="#5A5A5A")
    current_path = WORK / "diagram-kondisi-awal.png"
    current.save(current_path, quality=95)

    target = Image.new("RGB", (1800, 1020), "white")
    d = ImageDraw.Draw(target)
    d.text((70, 30), "ALUR PEMINJAMAN IPTOOLS LITE", font=bold, fill=f"#{NAVY}")
    lanes = [
        ("PEMOHON ATAU PIC AREA PEMINJAM", 125, "#EAF3FB"),
        ("REVIEWER AREA", 350, "#F5F1FB"),
        ("PIC AREA SUMBER DAN SISTEM", 575, "#EEF7EE"),
        ("PEMEGANG ALAT", 800, "#FFF5E8"),
    ]
    for label, y, fill in lanes:
        d.rounded_rectangle((50, y, 1750, y + 175), radius=16, fill=fill, outline="#B7C5D4", width=2)
        d.text((72, y + 14), label, font=small, fill=f"#{NAVY}")

    boxes = [
        (100, 180, "Login LDAP\nPilih alat dan periode", "#D9EAF7"),
        (445, 180, "Ajukan peminjaman\nIntra atau Antar Area", "#D9EAF7"),
        (210, 405, "Review area peminjam\nkhusus Antar Area", "#E6DCF5"),
        (600, 405, "Review area sumber\nsetujui jumlah atau tolak", "#E6DCF5"),
        (1030, 630, "Siapkan dan kirim\nfoto + Surat Jalan", "#DDEEDC"),
        (1390, 855, "Konfirmasi diterima", "#FFE6C5"),
        (920, 855, "Kembalikan sebagian\natau seluruhnya", "#FFE6C5"),
        (460, 630, "Verifikasi kembali\nSurat Jalan pengembalian", "#DDEEDC"),
        (115, 630, "Tutup transaksi\nstatus Selesai", "#DDEEDC"),
    ]
    centers = []
    for x, y, label, fill in boxes:
        box = (x, y, x + 285, y + 92)
        d.rounded_rectangle(box, radius=16, fill=fill, outline=f"#{DARK_BLUE}", width=3)
        draw_centered_text(d, box, label, small, "#111111")
        centers.append(((x, y), (x + 285, y + 92)))
    arrows = [
        ((385, 226), (445, 226)),
        ((590, 272), (352, 405)),
        ((495, 451), (600, 451)),
        ((885, 451), (1172, 630)),
        ((1315, 676), (1532, 855)),
        ((1390, 901), (1205, 901)),
        ((920, 901), (602, 722)),
        ((460, 676), (400, 676)),
    ]
    for a, b in arrows:
        arrow(d, a, b)
    d.text((935, 470), "Ditolak: alur berhenti dan alasan tersimpan", font=small, fill="#A61B1B")
    d.text((1150, 950), "Jika pengembalian baru sebagian, status tetap Dikembalikan Partials.", font=small, fill="#5A5A5A")
    target_path = WORK / "diagram-alur-sasaran.png"
    target.save(target_path, quality=95)
    return current_path, target_path


def patch_header_text(docx_path):
    replacements = {
        b"INFRASIGH": b"IPTOOLS LITE",
        b"INFRASIGHT": b"IPTOOLS LITE",
        b"InfraSight": b"IPTOOLS LITE",
    }
    fd, temp_name = tempfile.mkstemp(suffix=".docx", dir=str(docx_path.parent))
    os.close(fd)
    try:
        with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(temp_name, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.startswith("word/header") and item.filename.endswith(".xml"):
                    for old, new in replacements.items():
                        data = data.replace(old, new)
                    root = etree.fromstring(data)
                    namespaces = {"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
                    for doc_pr in root.xpath(".//wp:docPr", namespaces=namespaces):
                        if not doc_pr.get("descr"):
                            name = doc_pr.get("name", "")
                            alt = "Logo PLN Indonesia Power" if name.startswith("Picture") else "Elemen dekoratif header"
                            doc_pr.set("descr", alt)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
                zout.writestr(item, data)
        shutil.move(temp_name, docx_path)
    finally:
        if os.path.exists(temp_name):
            os.remove(temp_name)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    current_diagram, target_diagram = build_diagrams()
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    configure_styles(doc)

    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)

    section = doc.sections[0]
    section.different_first_page_header_footer = True

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    # Sampul
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    p.add_run("Business Requirements Document")
    p2 = add_body(doc, "BRD", bold=True, size=11, color=MID_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    p2.paragraph_format.keep_with_next = True
    banner = doc.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner.autofit = False
    banner.columns[0].width = Cm(18.0)
    set_repeat_table_header(banner.rows[0])
    cell = banner.cell(0, 0)
    set_cell_shading(cell, DARK_BLUE)
    set_cell_margins(cell, top=220, bottom=220, start=120, end=120)
    set_table_borders(banner, color=DARK_BLUE, size="4")
    bp = cell.paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    br = bp.add_run("IPTOOLS LITE")
    set_font(br, size=18, bold=True, color=WHITE)
    add_body(doc, "SISTEM PENGELOLAAN ALAT DAN PEMINJAMAN ANTAR AREA", bold=True, size=11, color=BLACK, align=WD_ALIGN_PARAGRAPH.CENTER, after=2).paragraph_format.space_before = Pt(15)
    add_body(doc, "MASTER ALAT, REVIEW, PENGIRIMAN, PENGEMBALIAN, MUTASI, LAPORAN, DAN AUDIT", size=10, color=MID_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    add_body(doc, "PT PLN INDONESIA POWER UBH", size=10, color=MID_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    add_body(doc, "VERSI 1.0  SEPTEMBER 2026", size=10, color=MID_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    doc.add_page_break()

    # Kontrol dokumen
    add_heading(doc, "KONTROL DOKUMEN", 1)
    add_heading(doc, "Riwayat Perubahan", 2)
    add_table(doc, ["Tanggal", "Penyusun", "Versi", "Uraian Perubahan"], [
        ["10/09/2026", "Business Analyst dan Tim Pengembang IPTOOLS LITE", "1.0", "Penyusunan BRD berdasarkan aplikasi, source code, manual, route, model, service, migration, dan pengujian otomatis."],
        ["", "", "", ""],
    ], [3.1, 4.8, 2.0, 8.1], 8.2)
    add_heading(doc, "Pemeriksaan dan Persetujuan Dokumen", 2)
    add_table(doc, ["Peran", "Nama atau Jabatan", "Status", "Tanggal dan Tanda Tangan"], [
        ["Disusun oleh", "Business Analyst dan Tim Pengembang IPTOOLS LITE", "Draft", ""],
        ["Diperiksa oleh", "Pemilik Proses Pengelolaan Tool  TBD", "Menunggu pemeriksaan", ""],
        ["Disetujui oleh", "Pemilik Proyek atau Pejabat Berwenang  TBD", "Menunggu persetujuan", ""],
    ], [3.2, 7.0, 2.7, 5.1], 8.2)
    add_heading(doc, "Distribusi Dokumen", 2)
    add_table(doc, ["Nomor Salinan", "Penerima atau Fungsi", "Lokasi atau Media"], [
        ["01", "Pemilik Proses dan Mgr Tool", "Repositori resmi proyek"],
        ["02", "SP Tool dan PIC Tool", "Salinan kerja terkendali"],
        ["03", "Tim Aplikasi, Infrastruktur, dan Database", "Salinan kerja terkendali"],
    ], [3.0, 8.3, 6.7], 8.2)
    add_note(doc, "Catatan pengendalian: nilai TBD memerlukan keputusan pemilik proses dan tidak dianggap sebagai komitmen yang telah disetujui.")

    # Daftar isi
    add_heading(doc, "DAFTAR ISI", 1, new_page=True)
    add_toc(doc)

    # Bab 1
    add_heading(doc, "1 DETAIL PROYEK RINGKASAN TUJUAN RUANG LINGKUP DAN PERAN", 1, new_page=True)
    add_heading(doc, "1.1 Detail Proyek", 2)
    add_table(doc, ["Atribut", "Isi"], [
        ["Nama Proyek", "IPTOOLS LITE"],
        ["Nama Modul atau Produk", "Dashboard; master alat; template peminjaman; peminjaman intra area dan antar area; review; pengiriman; penerimaan; pengembalian; mutasi; laporan kerusakan dan kehilangan; pengguna; area; log."],
        ["Pemilik Proyek", "TBD oleh organisasi"],
        ["Manajer Proyek", "TBD oleh pemilik proyek"],
        ["Pemilik Proses Bisnis", "Fungsi pengelolaan tool PT PLN Indonesia Power UBH  perlu konfirmasi"],
        ["Platform", "Laravel 10 dan PHP 8.2; Inertia.js 2; Vue 3; Tailwind CSS 4; Oracle sebagai koneksi basis data default."],
        ["Periode Sasaran", "Aplikasi berada dalam implementasi aktif pada 2026. Tanggal penerimaan formal dan go live ditetapkan setelah UAT."],
    ], [5.1, 12.9], 8.4)
    add_heading(doc, "1.2 Ringkasan Eksekutif dan Kebutuhan Bisnis", 2)
    add_body(doc, "IPTOOLS LITE adalah aplikasi web internal untuk mengendalikan data alat dan seluruh siklus peminjaman, mulai dari pemilihan berdasarkan ketersediaan, review per area, pengiriman dan penerimaan, pengembalian sebagian atau seluruhnya, sampai penutupan transaksi. Aplikasi juga menangani laporan kerusakan dan kehilangan, dokumen surat jalan, notifikasi, dashboard operasional, dan jejak aktivitas.")
    add_body(doc, "Sistem menggunakan akun lokal sebagai daftar pengguna yang diizinkan, sedangkan kredensial diverifikasi ke LDAP. Hak akses ditentukan oleh tujuh role dan konteks area. Transaksi penting memakai validasi server, transaksi basis data, row lock pada proses stok tertentu, serta pencatatan aktivitas agar perubahan dapat ditelusuri.")
    add_heading(doc, "1.3 Tujuan dan Indikator Keberhasilan", 2)
    add_table(doc, ["No", "Tujuan Bisnis", "Indikator atau Ukuran Keberhasilan", "Sasaran"], [
        ["1", "Menyediakan satu sumber data alat per area", "Kode alat unik per area, data dapat dicari, difilter, diimpor, dan diekspor", "Seluruh master aktif tersimpan dan dapat ditelusuri"],
        ["2", "Mengendalikan ketersediaan alat", "Permintaan tidak melebihi stok tersedia setelah memperhitungkan transaksi aktif", "Tidak ada alokasi melebihi stok pada skenario UAT"],
        ["3", "Menstandarkan peminjaman intra dan antar area", "Status dan review mengikuti kategori, area peminjam, dan area sumber", "Seluruh skenario Wajib lulus UAT"],
        ["4", "Menyediakan bukti serah terima", "Surat jalan dan foto tersimpan serta dapat diunduh oleh pihak berwenang", "Dokumen tersedia pada setiap pengiriman dan pengembalian"],
        ["5", "Meningkatkan keterlacakan", "Riwayat, log aktivitas, log alat, reviewer, jumlah, dan status dapat diaudit", "Sampel transaksi dapat ditelusuri dari awal sampai selesai"],
        ["6", "Mempercepat tindak lanjut", "Badge, mailbox, email status, dan pengingat jatuh tempo tersedia", "Pengguna menerima informasi sesuai perannya"],
    ], [1.0, 4.4, 7.2, 5.4], 7.8)
    add_heading(doc, "1.4 Ruang Lingkup dan Peran", 2)
    add_table(doc, ["Dalam Lingkup", "Di Luar Lingkup Versi Saat Ini", "Peran Utama"], [
        ["Login LDAP untuk akun yang diizinkan; logout; session berbasis web", "Registrasi mandiri, reset password lokal, dan pengelolaan password LDAP", "Semua role; Admin untuk allow list pengguna"],
        ["Master area dan alat, kode/lokasi/jenis/klasifikasi/stok; import dan export", "Pengadaan, depresiasi finansial, barcode, RFID, dan IoT", "Super Admin; PIC Tool; Admin; pembaca sesuai role"],
        ["Peminjaman intra area dan antar area, template, review per item", "Reservasi eksternal di luar jaringan atau integrasi vendor", "User; PIC Tool; SP Tool; Mgr Tool; Admin; Super Admin"],
        ["Pengiriman, penerimaan, pengembalian, surat jalan XLSX, dan foto", "Pelacakan ekspedisi melalui API kurir", "PIC Tool; peminjam; Admin; Super Admin"],
        ["Laporan kerusakan/kehilangan dan penyesuaian stok", "Work order perbaikan, klaim asuransi, dan proses penggantian aset", "SP Tool; PIC Tool; Mgr Tool; Admin; Super Admin"],
        ["Dashboard, notifikasi, email, log, export CSV, dan scheduler", "Data warehouse, aplikasi mobile native, dan analitik prediktif", "Semua role sesuai matriks akses; tim aplikasi"],
    ], [7.2, 6.4, 4.4], 7.9)

    # Bab 2
    add_heading(doc, "2 KAJIAN KELAYAKAN OPERASI", 1, new_page=True)
    add_heading(doc, "2.1 Permintaan dan Kebutuhan Perubahan", 2)
    add_body(doc, "Pengelolaan alat melibatkan banyak area, role, status, jumlah item, bukti foto, dan dokumen surat jalan. Tanpa satu alur terpusat, ketersediaan alat mudah berbeda antar catatan, persetujuan sulit dibuktikan, dan perpindahan alat antar area tidak segera tercermin pada stok operasional.")
    add_body(doc, "IPTOOLS LITE dibutuhkan untuk menempatkan validasi stok, kewenangan area, review, dokumen, dan audit trail dalam satu transaksi bisnis. Sistem juga harus menjaga alur tetap berjalan ketika pengiriman email gagal; kegagalan komunikasi dicatat untuk tindak lanjut operasional.")
    add_heading(doc, "2.2 Alternatif Solusi dan Kelayakan Operasi", 2)
    add_table(doc, ["Aspek", "Alternatif 1 Proses Manual", "Alternatif 2 IPTOOLS LITE", "Alternatif 3 Produk Eksternal"], [
        ["Proses bisnis", "Spreadsheet, pesan, dan dokumen terpisah", "Workflow terintegrasi dengan status dan area", "Konfigurasi produk dan migrasi proses"],
        ["Kontrol stok", "Rekonsiliasi manual", "Perhitungan availability, transaksi, dan row lock", "Tergantung modul vendor"],
        ["Hak akses", "Kontrol berbasis kepemilikan file", "Role dan area diverifikasi di route dan controller", "Tergantung lisensi dan konfigurasi"],
        ["Dokumen", "Surat jalan dibuat dan disimpan manual", "XLSX dibuat dari template bersama foto pengiriman/pengembalian", "Fitur standar mungkin perlu penyesuaian"],
        ["Integrasi", "Tidak terintegrasi", "LDAP, Oracle, SMTP, storage, queue, scheduler", "Perlu integrasi baru"],
        ["Biaya dan waktu", "Rendah di awal tetapi tinggi untuk koreksi", "Memanfaatkan source yang telah tersedia; perlu operasi dan UAT", "Pengadaan, konfigurasi, lisensi, dan migrasi"],
    ], [3.3, 4.7, 5.5, 4.5], 7.7)
    add_heading(doc, "2.3 Rekomendasi Kelayakan Operasi", 2)
    add_body(doc, "IPTOOLS LITE direkomendasikan sebagai solusi internal karena alur inti dan kontrol utama telah tersedia di source. Sebelum penerimaan formal, organisasi perlu menyetujui owner layanan, SLA, RPO, RTO, retensi dokumen/foto, konfigurasi queue dan scheduler, prosedur backup/restore, HTTPS, UAT, serta dukungan pasca implementasi.")

    # Bab 3
    add_heading(doc, "3 KAJIAN KELAYAKAN FINANSIAL RISIKO DAN KESIMPULAN", 1, new_page=True)
    add_heading(doc, "3.1 Analisis Biaya dan Manfaat", 2)
    add_table(doc, ["Kategori", "Komponen", "Estimasi atau Dampak", "Keterangan"], [
        ["Biaya", "Analisis, development, QA, security review, UAT, dokumentasi", "TBD", "Perlu estimasi resmi jam kerja dan perubahan tersisa"],
        ["Biaya", "Oracle, web server, storage surat jalan/foto, backup, SMTP, DNS/sertifikat", "TBD", "Gunakan kapasitas existing jika disetujui"],
        ["Biaya Berulang", "Operasi aplikasi, queue worker, scheduler, monitoring, patching, dukungan", "TBD per tahun", "Masukkan ke biaya layanan aplikasi"],
        ["Manfaat Tak Berwujud", "Waktu pencarian dan rekonsiliasi berkurang", "Perlu baseline", "Ukur waktu proses sebelum dan sesudah penggunaan"],
        ["Manfaat Tak Berwujud", "Kontrol stok, bukti transaksi, dan audit membaik", "Perlu sampling", "Ukur selisih stok dan temuan audit"],
    ], [2.8, 6.2, 3.6, 5.4], 8.0)
    add_heading(doc, "3.2 Identifikasi dan Analisis Risiko", 2)
    risks = [
        ["R-01", "Ketersediaan", "Aplikasi tidak dapat mengakses Oracle", "Gangguan jaringan atau koneksi idle", "Transaksi dan queue berhenti", "Health check, retry, worker dikelola Supervisor, monitoring koneksi"],
        ["R-02", "Autentikasi", "LDAP tidak tersedia", "Gangguan AD, DNS, CA, atau jaringan", "Pengguna gagal login", "Timeout, koordinasi tim AD, dan prosedur insiden"],
        ["R-03", "Otorisasi", "Role atau area salah", "Master pengguna tidak akurat", "Akses data yang tidak sesuai", "Review akses berkala, validasi backend, prinsip least privilege"],
        ["R-04", "Integritas stok", "Alokasi melebihi stok atau stok antar area salah", "Transaksi bersamaan atau data lama", "Operasi dan laporan tidak akurat", "Transaksi DB, row lock, constraint, rekonsiliasi berkala"],
        ["R-05", "Import", "File import mengubah data yang salah", "Format atau area tidak valid", "Master alat tercemar", "Template resmi, validasi seluruh baris, log hasil, pembatasan role"],
        ["R-06", "Dokumen", "Foto atau surat jalan hilang", "Storage penuh atau file terhapus", "Bukti serah terima tidak tersedia", "Storage privat, backup, monitoring kapasitas, uji download"],
        ["R-07", "Komunikasi", "Email terlambat atau gagal", "SMTP atau worker tidak aktif", "Review dan pengembalian terlambat", "Log kegagalan, badge dalam aplikasi, monitoring queue dan scheduler"],
        ["R-08", "Data laporan", "Status kerusakan/kehilangan mengubah stok tidak sesuai", "Transisi status salah atau proses ganda", "Total aset keliru", "Row lock, UAT transisi ulang, audit log, rekonsiliasi"],
        ["R-09", "Privasi", "Foto atau data pegawai diakses pihak tak berwenang", "URL atau storage salah konfigurasi", "Kebocoran data", "Route terautorisasi, retensi, pembatasan akses, audit"],
        ["R-10", "Transport", "Akses melalui HTTP", "HTTPS/DNS belum final", "Credential dapat disadap", "Wajib HTTPS dan sertifikat tepercaya sebelum perluasan akses"],
        ["R-11", "Operasi", "Scheduler tidak berjalan", "Task scheduler belum dikonfigurasi", "LDAP dan reminder tidak diperbarui", "Monitoring job 01.00 dan 07.00, runbook, alert kegagalan"],
        ["R-12", "Perubahan", "Release menyebabkan regresi", "Tidak ada gate QA/UAT/rollback", "Downtime atau data rusak", "CI/test, backup pra-rilis, migration review, smoke test, rollback"],
    ]
    add_table(doc, ["ID", "Sasaran", "Kejadian Risiko", "Penyebab", "Dampak", "Mitigasi"], risks, [1.0, 2.3, 3.7, 3.4, 3.3, 4.3], 6.8)
    add_heading(doc, "3.3 Kesimpulan Kajian Kelayakan", 2)
    add_body(doc, "Inisiatif layak dilanjutkan. Source saat ini memiliki alur inti dan test otomatis yang mencakup fungsi penting, tetapi kelayakan produksi tetap bergantung pada keputusan operasional, uji keamanan, data awal, UAT bisnis, dan kesiapan infrastruktur.")

    # Bab 4
    add_heading(doc, "4 DESAIN PROSES KONDISI SAAT INI", 1, new_page=True)
    add_note(doc, "Bagian ini menggambarkan proses pembanding sebelum aplikasi atau proses manual yang hendak digantikan. Pemilik proses perlu memvalidasi detailnya.")
    add_heading(doc, "4.1 Gambaran Proses Saat Ini", 2)
    add_table(doc, ["Tahap", "Aktor atau Fungsi", "Aktivitas", "Masukan", "Keluaran", "Kendala"], [
        ["1", "Pemohon", "Menyampaikan kebutuhan alat dan periode", "Kebutuhan pekerjaan", "Permintaan", "Format tidak selalu seragam"],
        ["2", "Pengelola alat", "Memeriksa stok pada catatan area", "Master atau spreadsheet", "Konfirmasi stok", "Data dapat berbeda antar sumber"],
        ["3", "Atasan atau reviewer", "Memberi persetujuan melalui komunikasi terpisah", "Permintaan", "Persetujuan", "Alasan dan waktu sulit ditelusuri"],
        ["4", "PIC area", "Menyerahkan atau mengirim alat serta bukti", "Daftar alat", "Surat jalan dan foto", "Bukti tersebar"],
        ["5", "Pemegang alat", "Mengembalikan alat", "Alat dan bukti", "Konfirmasi kembali", "Pengembalian sebagian sulit direkap"],
        ["6", "Pengelola", "Merekap mutasi, kerusakan, kehilangan, dan stok", "Dokumen terpisah", "Laporan", "Rekonsiliasi lambat"],
    ], [1.1, 2.7, 4.3, 3.0, 3.0, 3.9], 7.1)
    add_heading(doc, "4.2 Diagram Proses Saat Ini", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    current_picture = p.add_run().add_picture(str(current_diagram), width=Inches(7.0))
    current_picture._inline.docPr.set("descr", "Diagram kondisi proses pembanding sebelum IPTOOLS LITE")

    # Bab 5
    add_heading(doc, "5 DESAIN PROSES KONDISI SASARAN", 1, new_page=True)
    add_heading(doc, "5.1 Gambaran Proses Sasaran", 2)
    add_table(doc, ["Tahap", "Aktor atau Fungsi", "Aktivitas Sasaran", "Sistem atau Media", "Keluaran", "Kontrol"], [
        ["1", "Pengguna", "Login dengan username dan password LDAP", "IPTOOLS LITE dan LDAP", "Session sesuai role", "Akun harus ada pada allow list lokal"],
        ["2", "Pemohon atau PIC", "Pilih alat dan periode berdasarkan availability", "Katalog alat dan template", "Draft permintaan", "Area dan stok divalidasi server"],
        ["3", "Sistem", "Menyimpan pengajuan dan menentukan tahap review", "Oracle", "Perlu Disetujui atau Perlu Direview", "Kategori intra atau antar area"],
        ["4", "SP Tool atau Mgr Tool", "Review seluruh item dan jumlah", "Review Peminjaman", "Disetujui, ditolak, atau lanjut review", "Area reviewer dan jumlah maksimum"],
        ["5", "PIC area sumber", "Siapkan, foto, buat Surat Jalan, dan kirim", "Storage privat dan XLSX", "Dikirim", "Minimal satu foto, maksimal delapan"],
        ["6", "Peminjam atau PIC area peminjam", "Konfirmasi penerimaan", "Mutasi atau transaksi", "Diterima", "Antar area membentuk stok sementara"],
        ["7", "Pemegang alat", "Kembalikan sebagian atau seluruh alat dan bukti", "Form pengembalian", "Dikembalikan Partials atau Semuanya", "Jumlah tidak melebihi sisa"],
        ["8", "PIC area sumber", "Verifikasi dan tutup transaksi", "Transaksi", "Selesai", "Hanya setelah seluruh alat kembali"],
    ], [1.0, 2.8, 4.7, 3.5, 3.2, 4.0], 6.9)
    add_heading(doc, "5.2 Diagram Proses Sasaran", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    target_picture = p.add_run().add_picture(str(target_diagram), width=Inches(7.05))
    target_picture._inline.docPr.set("descr", "Diagram alur proses sasaran peminjaman IPTOOLS LITE")

    # Bab 6
    add_heading(doc, "6 ALUR PROSES BISNIS", 1, new_page=True)
    add_heading(doc, "6.1 Alur Peminjaman Intra Area", 2)
    add_table(doc, ["Urutan", "Aktor", "Kegiatan", "Status Hasil"], [
        ["1", "User, Admin, atau Super Admin", "Mengisi periode, pekerjaan, area, dan item dengan jumlah lebih dari nol", "Perlu Disetujui"],
        ["2", "SP Tool, Mgr Tool, atau Super Admin", "Review semua item; dapat menyetujui jumlah parsial atau menolak dengan alasan", "Disetujui atau Ditolak"],
        ["3", "PIC Tool, Admin, atau Super Admin", "Mengisi pengirim, resi opsional, dan 1 sampai 8 foto; sistem membuat Surat Jalan", "Dikirim"],
        ["4", "Peminjam, PIC Tool, Admin, atau Super Admin", "Mengonfirmasi alat diterima sesuai kewenangan", "Diterima"],
        ["5", "Peminjam, Admin, atau Super Admin", "Mengembalikan item dan foto; dapat sekaligus melaporkan kerusakan/kehilangan", "Dikembalikan Partials atau Dikembalikan Semuanya"],
        ["6", "PIC Tool, Admin, atau Super Admin", "Menutup transaksi setelah seluruh alat kembali", "Selesai"],
    ], [1.2, 3.4, 9.6, 3.8], 7.6)
    add_heading(doc, "6.2 Alur Peminjaman Antar Area", 2)
    add_table(doc, ["Urutan", "Aktor", "Kegiatan", "Status Hasil"], [
        ["1", "PIC Tool, Admin, atau Super Admin", "Memilih area sumber yang berbeda dari area peminjam, periode, dan alat tersedia", "Perlu Direview"],
        ["2", "Reviewer area peminjam", "Menilai kebutuhan dan jumlah semua item", "Perlu Disetujui atau Ditolak"],
        ["3", "Reviewer area sumber", "Menilai kembali jumlah tanpa melebihi hasil review pertama", "Disetujui atau Ditolak"],
        ["4", "PIC area sumber", "Mengirim alat dengan foto dan Surat Jalan", "Dikirim"],
        ["5", "PIC area peminjam", "Mengonfirmasi penerimaan; sistem mencatat stok sementara berdasarkan transaksi sumber", "Diterima"],
        ["6", "PIC area peminjam", "Mengembalikan sebagian atau seluruh alat dengan foto dan dokumen", "Dikembalikan Partials atau Semuanya"],
        ["7", "PIC area sumber", "Menutup transaksi setelah seluruh alat diterima kembali", "Selesai"],
    ], [1.2, 3.5, 9.5, 3.8], 7.6)
    add_heading(doc, "6.3 Aturan Bisnis pada Alur", 2)
    business_rules = [
        ["AB-01", "Hanya akun lokal yang diizinkan dan kredensial LDAP yang valid dapat login.", "Login", "Session dibuat dan diregenerasi; akses lain ditolak."],
        ["AB-02", "Role dan area diperiksa pada route dan controller.", "Setiap aksi", "Data dan aksi dibatasi sesuai konteks pengguna."],
        ["AB-03", "Tanggal kembali harus sama atau setelah tanggal pinjam.", "Pengajuan atau ubah periode", "Input tidak valid ditolak."],
        ["AB-04", "Jumlah diminta minimal satu dan tidak boleh melebihi availability.", "Simpan peminjaman", "Stok diperiksa di dalam transaksi."],
        ["AB-05", "Semua item harus direview; jumlah disetujui tidak melebihi jumlah yang tersedia pada tahap sebelumnya.", "Review", "Status item dan alasan penolakan disimpan."],
        ["AB-06", "Peminjaman antar area melalui review area peminjam lalu area sumber.", "Kategori Antar Area", "Status berpindah dari Perlu Direview ke Perlu Disetujui."],
        ["AB-07", "Pengiriman hanya dari status Disetujui dan penerimaan hanya dari status Dikirim.", "Aksi kirim atau terima", "Transisi di luar urutan ditolak."],
        ["AB-08", "Pengiriman dan pengembalian memerlukan 1 sampai 8 gambar JPG, PNG, atau WebP maksimal 5 MB dan 6000 x 6000 piksel per gambar.", "Unggah bukti", "Gambar dikompresi; Surat Jalan XLSX dibuat."],
        ["AB-09", "Pengembalian tidak boleh melebihi jumlah yang belum kembali.", "Kembalikan", "Status parsial dipertahankan sampai seluruh jumlah kembali."],
        ["AB-10", "Antar area membentuk stok sementara saat diterima dan menguranginya saat dikembalikan.", "Terima atau kembali", "AreaAlatStock mengikuti sumber peminjaman."],
        ["AB-11", "Kode alat unik dalam satu area; jika kosong, sistem membuat kode dari awalan area dan nomor urut.", "Tambah atau ubah alat", "Duplikasi kode dalam area ditolak."],
        ["AB-12", "Import alat menerima CSV, TXT, atau XLSX maksimal 5 MB dan diproses sebagai job.", "Import", "Status pending, processing, completed, atau failed tersedia."],
        ["AB-13", "Laporan kerusakan/kehilangan memerlukan jumlah, deskripsi maksimal 1000 karakter, dan satu foto maksimal 5 MB.", "Buat laporan", "Status awal Dilaporkan."],
        ["AB-14", "Perubahan laporan ke Disetujui mengurangi stok; keluar dari Disetujui mengembalikan stok sesuai konteks area.", "Ubah status laporan", "Penyesuaian dilakukan dalam transaksi dan row lock."],
        ["AB-15", "Bulk delete tertentu memerlukan verifikasi password LDAP dan kewenangan admin.", "Hapus data area", "Aksi berisiko tidak dapat dijalankan tanpa autentikasi ulang."],
    ]
    add_table(doc, ["ID", "Aturan Bisnis", "Pemicu atau Kondisi", "Aksi atau Hasil"], business_rules, [1.2, 8.0, 4.0, 4.8], 7.1)

    # Bab 7
    add_heading(doc, "7 DESAIN STATUS MENU DAN HAK AKSES", 1, new_page=True)
    add_heading(doc, "7.1 Desain Status Proses", 2)
    status_rows = [
        ["Perlu Direview", "Antar area menunggu review area peminjam", "Pengajuan antar area", "Review semua item", "Perlu Disetujui atau Ditolak"],
        ["Perlu Disetujui", "Menunggu review area sumber atau reviewer intra area", "Pengajuan intra atau review pertama", "Setujui jumlah atau tolak", "Disetujui atau Ditolak"],
        ["Disetujui", "Ada item disetujui dan siap dikirim", "Review final", "Kirim", "Dikirim"],
        ["Ditolak", "Tidak ada item yang diteruskan", "Review", "Lihat riwayat", "Final"],
        ["Dikirim", "Alat dan Surat Jalan telah dikirim", "Kirim", "Terima", "Diterima"],
        ["Diterima", "Alat berada pada pemegang atau area peminjam", "Konfirmasi penerimaan", "Ubah periode intra area atau kembalikan", "Dikembalikan Partials atau Semuanya"],
        ["Dikembalikan Partials", "Sebagian alat telah dikembalikan", "Pengembalian", "Kembalikan sisa", "Dikembalikan Semuanya"],
        ["Dikembalikan Semuanya", "Semua jumlah disetujui telah kembali", "Pengembalian final", "Selesaikan", "Selesai"],
        ["Selesai", "Transaksi ditutup", "Verifikasi PIC", "Lihat riwayat dan dokumen", "Final"],
    ]
    add_table(doc, ["Status", "Definisi", "Pemicu Masuk", "Aksi yang Diizinkan", "Status Berikutnya"], status_rows, [3.1, 5.0, 3.2, 4.0, 3.7], 7.3)
    add_heading(doc, "7.2 Status Pendukung", 2)
    add_table(doc, ["Objek", "Status", "Makna"], [
        ["Item review", "Menunggu Review; Disetujui; Ditolak", "Keputusan dan jumlah pada tingkat item."],
        ["Laporan alat", "Dilaporkan; Disetujui; Ditolak; Selesai", "Siklus laporan kerusakan atau kehilangan dan dampak stok."],
        ["Import alat", "pending; processing; completed; failed", "Status pekerjaan import dan hasil pemrosesan."],
        ["Surat Jalan", "peminjaman; pengembalian", "Jenis dokumen dengan urutan untuk pengembalian berulang."],
        ["Stok area", "active true atau false", "Ketersediaan alat yang berasal dari transaksi antar area."],
    ], [3.2, 6.2, 8.8], 8.1)
    add_heading(doc, "7.3 Desain Menu dan Hak Akses", 2)
    menu_rows = [
        ["Dashboard", "Semua role", "Ringkasan area; insight operasional untuk role tertentu"],
        ["Buat Peminjaman", "User, Admin, Super Admin", "Buat peminjaman intra area"],
        ["Peminjaman Antar Area", "PIC Tool, Admin, Super Admin", "Buat permintaan antar area"],
        ["Review Peminjaman", "SP Tool, Mgr Tool, Super Admin", "Review sesuai area dan tahap"],
        ["Alat", "Guest, SP Tool, PIC Tool, Mgr Tool, Admin, Super Admin", "Baca; CRUD/import khusus PIC Tool, Admin, Super Admin"],
        ["Template Peminjaman", "PIC Tool, Admin, Super Admin", "Buat, ubah, hapus template; template dapat dibaca role terautentikasi"],
        ["Intra Area dan Antar Area", "PIC Tool, Admin, Super Admin", "Pengiriman, penerimaan, pengembalian, periode, selesai"],
        ["Mutasi Alat", "Guest, User, Mgr Tool, Super Admin", "Pantau mutasi; User melakukan aksi atas transaksi miliknya"],
        ["Riwayat Peminjaman", "Semua role", "Lihat/filter/export; edit/hapus dibatasi Admin dan Super Admin"],
        ["Laporan Kerusakan dan Kehilangan", "Guest, SP Tool, PIC Tool, Mgr Tool, Admin, Super Admin", "Guest hanya baca; role operasional dapat membuat dan mengubah status"],
        ["Pengguna", "SP Tool, Mgr Tool, Admin, Super Admin", "SP/Mgr melihat; LDAP lookup dan CRUD hanya Admin/Super Admin"],
        ["Area", "Super Admin", "Kelola master area"],
        ["Log Activity", "Guest, SP Tool, Mgr Tool, Admin, Super Admin", "Lihat dan export sesuai cakupan"],
        ["Log Alat", "Guest, SP Tool, PIC Tool, Mgr Tool, Admin, Super Admin", "Lihat dan export perubahan alat"],
    ]
    add_table(doc, ["Menu", "Role yang Dapat Mengakses", "Kewenangan Utama"], menu_rows, [4.2, 7.1, 6.9], 7.5)
    add_heading(doc, "7.4 Notifikasi dan Eskalasi", 2)
    notif_rows = [
        ["Pengajuan perlu review", "SP Tool pada area yang sedang mereview", "Email dan badge", "Setelah pengajuan atau review tahap pertama"],
        ["Peminjaman siap dikirim", "PIC Tool area sumber", "Email dan badge", "Setelah review final menghasilkan Disetujui"],
        ["Alat dikirim", "Peminjam intra area atau PIC Tool area peminjam", "Email dan mailbox", "Setelah status Dikirim"],
        ["Pengingat pengembalian", "Pemegang alat", "Email", "H-1, hari tenggat, lalu setiap 6 hari setelah terlambat"],
        ["Alat dikembalikan", "PIC Tool area sumber", "Email dan badge", "Setelah pengembalian parsial atau seluruhnya"],
        ["Transaksi selesai", "Peminjam awal", "Email", "Setelah status Selesai"],
        ["Laporan pending", "Role operasional sesuai area", "Badge dan mailbox", "Selama status Dilaporkan"],
        ["Email gagal", "Tim aplikasi", "Log aplikasi", "Pada exception; workflow inti tidak dibatalkan"],
    ]
    add_table(doc, ["Peristiwa", "Penerima", "Kanal", "Waktu atau Kondisi"], notif_rows, [4.4, 5.5, 3.1, 5.2], 7.8)

    # Bab 8
    add_heading(doc, "8 KEBUTUHAN BISNIS DAN KEBUTUHAN FUNGSIONAL", 1, new_page=True)
    add_heading(doc, "8.1 Kebutuhan Bisnis", 2)
    kb_rows = [
        ["KB-01", "Satu sumber data alat dan stok per area", "Pemilik data alat", "Wajib", "Kode unik, area, lokasi, jenis, klasifikasi, total, dipinjam, dan tersedia dapat diperiksa"],
        ["KB-02", "Siklus peminjaman intra area terkontrol", "Pemilik proses", "Wajib", "Pengajuan sampai selesai mengikuti status dan actor yang benar"],
        ["KB-03", "Siklus peminjaman antar area terkontrol", "Pemilik proses", "Wajib", "Dua tahap review dan stok area sementara dapat ditelusuri"],
        ["KB-04", "Pemisahan akses berdasarkan role dan area", "Pemilik aplikasi", "Wajib", "Menu dan API menolak akses di luar kewenangan"],
        ["KB-05", "Bukti pengiriman dan pengembalian tersimpan", "PIC Tool", "Wajib", "Foto dan Surat Jalan dapat diunduh oleh pihak berwenang"],
        ["KB-06", "Kerusakan dan kehilangan terhubung ke stok", "Pengelola alat", "Wajib", "Transisi status menghasilkan penyesuaian stok yang konsisten"],
        ["KB-07", "Informasi tindak lanjut tersedia tepat waktu", "Operasional", "Penting", "Badge, mailbox, email, dan reminder sesuai kejadian"],
        ["KB-08", "Data dapat dimuat dan dikeluarkan secara terkendali", "Admin data", "Penting", "Import tervalidasi serta export mengikuti filter dan area"],
        ["KB-09", "Aktivitas dan perubahan alat dapat diaudit", "Audit dan manajemen", "Wajib", "Log menyimpan actor, area, aksi, nilai lama/baru, route, IP, dan waktu"],
        ["KB-10", "Solusi dapat dioperasikan dan dipulihkan", "Tim aplikasi/infrastruktur", "Wajib", "Queue, scheduler, backup, monitoring, dan prosedur restore disetujui"],
    ]
    add_table(doc, ["ID", "Kebutuhan Bisnis", "Sumber atau Pemilik", "Prioritas", "Kriteria Keberhasilan"], kb_rows, [1.3, 5.4, 3.8, 2.2, 5.3], 7.3)
    add_heading(doc, "8.2 Kebutuhan Fungsional", 2)
    kf_rows = [
        ["KF-01", "Login dan Logout LDAP", "Validasi username/password ke LDAP untuk akun lokal yang diizinkan; regenerasi session dan invalidasi logout.", "Semua role", "Wajib", "KB-04"],
        ["KF-02", "Sinkronisasi LDAP", "Sinkronkan user direktori ke cache lokal setiap hari pukul 01.00 tanpa tumpang tindih.", "Sistem", "Penting", "KB-04"],
        ["KF-03", "Manajemen Pengguna", "Cari user LDAP; tetapkan role dan area; ubah atau hapus sesuai kewenangan Admin/Super Admin.", "Admin, Super Admin", "Wajib", "KB-04"],
        ["KF-04", "Manajemen Area", "Tambah, ubah, hapus, cari, dan paginasi area dengan nama, slug, dan kode.", "Super Admin", "Wajib", "KB-01"],
        ["KF-05", "Dashboard Area", "Tampilkan total peminjaman, total aset/jenis, review, transaksi berjalan, selesai, dan laporan aktif.", "Semua role", "Penting", "KB-01, KB-07"],
        ["KF-06", "Insight Operasional", "Tampilkan kerusakan dan kehilangan per bulan serta total lintas area untuk role tertentu.", "Guest, SP, PIC, Mgr, Super Admin", "Penting", "KB-07"],
        ["KF-07", "Daftar Alat", "Cari, filter area, paginasi, dan hitung total, dipinjam, tersedia, serta sumber stok antar area.", "Role terautentikasi sesuai route", "Wajib", "KB-01"],
        ["KF-08", "CRUD Alat", "Kelola kode, nama, jenis, klasifikasi, lokasi, total aset, dan area.", "PIC Tool, Admin, Super Admin", "Wajib", "KB-01"],
        ["KF-09", "Import Alat", "Unduh template; unggah CSV/TXT/XLSX maksimal 5 MB; proses job dan tampilkan hasil.", "PIC Tool, Admin, Super Admin", "Penting", "KB-08"],
        ["KF-10", "Export Alat", "Export CSV berdasarkan pencarian, filter, dan cakupan area.", "Role pembaca alat", "Penting", "KB-08"],
        ["KF-11", "Template Peminjaman", "Kelola template item dan jumlah untuk intra area atau antar area.", "PIC Tool, Admin, Super Admin", "Penting", "KB-02, KB-03"],
        ["KF-12", "Buat Peminjaman Intra", "Validasi periode, pekerjaan, area, item, availability, lalu simpan Perlu Disetujui.", "User, Admin, Super Admin", "Wajib", "KB-02"],
        ["KF-13", "Buat Peminjaman Antar Area", "Pilih area sumber berbeda, cek stok, lalu simpan Perlu Direview.", "PIC Tool, Admin, Super Admin", "Wajib", "KB-03"],
        ["KF-14", "Review Area Peminjam", "Review semua item antar area dan batasi jumlah yang diteruskan.", "SP Tool, Mgr Tool, Super Admin", "Wajib", "KB-03"],
        ["KF-15", "Review Area Sumber", "Setujui atau tolak setiap item dengan jumlah dan alasan; kirim notifikasi siap kirim.", "SP Tool, Mgr Tool, Super Admin", "Wajib", "KB-02, KB-03"],
        ["KF-16", "Riwayat dan Export Peminjaman", "Cari, filter, lihat detail, dokumen, dan export CSV.", "Semua role", "Wajib", "KB-02, KB-03, KB-08"],
        ["KF-17", "Pengiriman", "Validasi status, pengirim, resi, dan foto; kompres foto dan buat Surat Jalan XLSX.", "PIC Tool, Admin, Super Admin", "Wajib", "KB-05"],
        ["KF-18", "Penerimaan", "Konfirmasi penerimaan oleh peminjam intra atau PIC area peminjam antar area.", "User, PIC Tool, Admin, Super Admin", "Wajib", "KB-02, KB-03"],
        ["KF-19", "Pengembalian", "Terima jumlah parsial/seluruh, foto, laporan opsional, dan buat Surat Jalan pengembalian.", "User, PIC requester, Admin, Super Admin", "Wajib", "KB-05, KB-06"],
        ["KF-20", "Ubah Periode", "Ubah tanggal kembali untuk peminjaman intra area berstatus Diterima.", "PIC Tool, Admin, Super Admin", "Penting", "KB-02"],
        ["KF-21", "Penyelesaian", "Tutup transaksi hanya dari Dikembalikan Semuanya.", "PIC Tool, Admin, Super Admin", "Wajib", "KB-02, KB-03"],
        ["KF-22", "Mutasi Alat", "Tampilkan riwayat pengiriman, penerimaan, pengembalian, item, dan Surat Jalan.", "Role sesuai menu", "Wajib", "KB-05, KB-09"],
        ["KF-23", "Laporan Kerusakan", "Buat laporan dengan alat, jumlah, deskripsi, foto, dan area.", "SP, PIC, Mgr, Admin, Super Admin", "Wajib", "KB-06"],
        ["KF-24", "Laporan Kehilangan", "Buat laporan dengan alat, jumlah, deskripsi, foto, dan area.", "SP, PIC, Mgr, Admin, Super Admin", "Wajib", "KB-06"],
        ["KF-25", "Status Laporan", "Ubah status ke Disetujui, Ditolak, atau Selesai serta sesuaikan stok atomik.", "SP, PIC, Mgr, Admin, Super Admin", "Wajib", "KB-06"],
        ["KF-26", "Export dan Template Laporan", "Unduh template surat dan export CSV berdasarkan filter.", "Role pembaca laporan", "Penting", "KB-08"],
        ["KF-27", "Notifikasi Dalam Aplikasi", "Tampilkan badge dan mailbox untuk review, kirim, mutasi, kerusakan, dan kehilangan.", "Sesuai role dan area", "Penting", "KB-07"],
        ["KF-28", "Email Transaksional", "Kirim email review, siap kirim, dikirim, dikembalikan, selesai, dan reminder.", "Sistem", "Penting", "KB-07"],
        ["KF-29", "Pengingat Pengembalian", "Jalankan pukul 07.00 untuk H-1, hari H, dan setiap enam hari keterlambatan.", "Sistem", "Penting", "KB-07"],
        ["KF-30", "Activity Log", "Cari, filter, paginasi, dan export aktivitas menurut kewenangan.", "Role audit sesuai route", "Wajib", "KB-09"],
        ["KF-31", "Log Alat", "Tampilkan dan export perubahan master/stok alat beserta delta dan actor.", "Role pembaca log alat", "Wajib", "KB-09"],
        ["KF-32", "Download Dokumen", "Izinkan download Surat Jalan hanya untuk user/role/area yang terkait.", "Pihak berwenang", "Wajib", "KB-05, KB-04"],
    ]
    add_table(doc, ["ID", "Fitur atau Fungsi", "Deskripsi", "Aktor", "Prioritas", "Referensi KB"], kf_rows, [1.2, 3.1, 7.4, 3.1, 2.0, 2.2], 6.5)
    add_note(doc, "Prioritas memakai MoSCoW: Wajib setara Must, Penting setara Should. Tidak ada fitur Diinginkan yang dinyatakan sebagai komitmen pada versi ini.")

    # Bab 9
    add_heading(doc, "9 KEBUTUHAN NONFUNGSIONAL DATA INTEGRASI DAN PELAPORAN", 1, new_page=True)
    add_heading(doc, "9.1 Kebutuhan Nonfungsional", 2)
    nfr_rows = [
        ["Kinerja", "Pencarian, pagination, availability, import, export, dan dashboard tetap responsif pada volume produksi.", "Uji beban dengan volume data representatif; target angka TBD."],
        ["Ketersediaan", "Web, Oracle, storage, worker, scheduler, LDAP, dan SMTP dipantau sebagai satu layanan.", "Health check, restart test, dan monitoring; SLA TBD."],
        ["Keamanan", "LDAP allow list, session regeneration, CSRF, RBAC, area scoping, validasi file, dan route terautorisasi.", "Security review, test akses silang role/area, dependency audit."],
        ["Privasi", "Atribut pegawai, foto, alamat IP, user agent, dan dokumen hanya dipakai untuk proses dan audit.", "Review data model, akses storage, retensi, dan masking bila diperlukan."],
        ["Integritas", "Stok dan status diperbarui atomik pada operasi kritis; jumlah tidak negatif atau melebihi sisa.", "Concurrency test, constraint review, dan rekonsiliasi sampel."],
        ["Kompatibilitas", "Antarmuka Vue/Inertia berjalan pada browser desktop modern yang disetujui organisasi.", "UAT pada Chrome/Edge versi organisasi; matriks resmi TBD."],
        ["Auditabilitas", "Actor, role, area, route, metode, IP, nilai lama/baru, dan waktu tersedia untuk transaksi penting.", "Sampling Activity Log dan Log Alat."],
        ["Pemulihan", "Database, dokumen, dan foto dapat dipulihkan sesuai RPO/RTO.", "Backup terjadwal dan drill restore; target RPO/RTO TBD."],
        ["Maintainability", "Source berada dalam Git, memakai migration, test otomatis, konfigurasi environment, dan dependency lock.", "Review repository, pipeline, dan prosedur release."],
        ["Operabilitas", "Queue worker memakai Supervisor atau service manager; scheduler aktif; kegagalan tercatat.", "Reboot test, stop-worker test, retry job, dan review log."],
        ["Transport", "Akses produksi memakai DNS resmi dan HTTPS dengan sertifikat tepercaya.", "TLS scan dan validasi browser sebelum akses diperluas."],
    ]
    add_table(doc, ["Kategori", "Kebutuhan atau Sasaran", "Metode Verifikasi"], nfr_rows, [3.2, 10.0, 4.8], 7.4)
    add_heading(doc, "9.2 Data dan Integrasi", 2)
    integration_rows = [
        ["Identitas pengguna", "Active Directory/LDAP", "Login dan sinkronisasi cache user", "LDAPRecord bind dan query; jadwal 01.00", "Real time saat login; harian untuk sinkronisasi", "Tim AD dan pemilik aplikasi"],
        ["Data aplikasi", "Oracle", "Master, transaksi, log, jobs, dan status", "Eloquent/OCI8, migration, transaksi", "Per transaksi", "DBA dan tim aplikasi"],
        ["Antarmuka", "Laravel Inertia dan Vue", "Menyajikan page dan memanggil API", "Session web, Axios, JSON", "Saat akses", "Tim aplikasi"],
        ["Email", "SMTP", "Review, pengiriman, pengembalian, selesai, reminder", "Laravel Mail; queue sesuai environment", "Per kejadian dan 07.00", "Tim email dan aplikasi"],
        ["Dokumen", "Storage lokal/private dan public terbatas", "Surat Jalan XLSX, foto, template laporan", "Filesystem Laravel dan route download", "Per transaksi", "Tim aplikasi/infrastruktur"],
        ["Import/Export", "CSV, TXT, XLSX", "Pertukaran master alat dan laporan", "PhpSpreadsheet, streamed CSV, job import", "Ad hoc", "Admin data"],
        ["Scheduler", "Laravel scheduler", "LDAP sync dan reminder pengembalian", "Task Scheduler atau cron", "01.00 dan 07.00", "Administrator server"],
        ["Audit", "Activity Log", "Jejak request dan perubahan model", "Observer dan ActivityLogger", "Per kejadian", "Pemilik aplikasi/audit"],
    ]
    add_table(doc, ["Objek atau Sistem", "Sumber", "Tujuan", "Metode", "Frekuensi", "Pemilik"], integration_rows, [3.1, 3.4, 4.4, 4.0, 3.0, 3.1], 6.7)
    add_heading(doc, "9.3 Pelaporan Dashboard dan Keluaran", 2)
    report_rows = [
        ["LP-01", "Dashboard Area", "Semua role", "Total peminjaman, aset, jenis alat, review, berjalan, selesai, laporan aktif", "Saat akses", "Web"],
        ["LP-02", "Insight Tahunan", "Guest, SP, PIC, Mgr, Super Admin", "Kerusakan dan kehilangan per bulan serta total aset lintas area", "Saat akses", "Web"],
        ["LP-03", "Master Alat", "Role pembaca alat", "Kode, nama, jenis, klasifikasi, lokasi, area, total, dipinjam, tersedia", "Ad hoc", "Web/CSV"],
        ["LP-04", "Riwayat Peminjaman", "Semua role", "Peminjam, area, periode, item, jumlah, review, status, dokumen", "Ad hoc", "Web/CSV"],
        ["LP-05", "Mutasi Alat", "Role sesuai menu", "Pengiriman, penerimaan, pengembalian, sisa, Surat Jalan", "Ad hoc", "Web/XLSX download"],
        ["LP-06", "Laporan Kerusakan", "Role laporan", "Alat, jumlah, deskripsi, foto, area, pelapor, status", "Ad hoc", "Web/CSV/DOCX template"],
        ["LP-07", "Laporan Kehilangan", "Role laporan", "Alat, jumlah, deskripsi, foto, area, pelapor, status", "Ad hoc", "Web/CSV/DOCX template"],
        ["LP-08", "Activity Log", "Role audit", "Actor, role, area, aksi, subject, route, IP, perubahan, waktu", "Ad hoc", "Web/CSV"],
        ["LP-09", "Log Alat", "Role pembaca log", "Perubahan alat, stok, delta, import, actor, waktu", "Ad hoc", "Web/CSV"],
        ["LP-10", "Status Import", "Pengunggah atau admin berwenang", "Jumlah diproses, dibuat, diubah, pesan gagal", "Saat import", "Web/file sumber"],
    ]
    add_table(doc, ["ID", "Laporan atau Dashboard", "Pengguna", "Isi atau KPI", "Frekuensi", "Format"], report_rows, [1.2, 3.4, 3.8, 7.0, 2.6, 2.4], 6.6)

    # Bab 10
    add_heading(doc, "10 DESAIN ANTARMUKA", 1, new_page=True)
    add_note(doc, "Daftar layar berikut merepresentasikan fungsi yang ada pada source. Desain visual final mengikuti komponen Vue yang berlaku dan hasil UAT.")
    add_heading(doc, "10.1 Daftar Layar dan Halaman", 2)
    ui_rows = [
        ["UI-01", "Login", "Masuk memakai username/password LDAP", "Pengguna belum login", "KF-01"],
        ["UI-02", "Dashboard", "Ringkasan area, kartu status, dan grafik insight", "Semua role", "KF-05, KF-06"],
        ["UI-03", "Buat Peminjaman", "Cari alat, keranjang, template, periode, dan checkout", "User, Admin, Super Admin", "KF-11, KF-12"],
        ["UI-04", "Peminjaman Antar Area", "Pilih area sumber, alat, template, periode, dan kirim", "PIC, Admin, Super Admin", "KF-11, KF-13"],
        ["UI-05", "Review Peminjaman", "Daftar pending, detail, keputusan item, jumlah, alasan", "SP, Mgr, Super Admin", "KF-14, KF-15"],
        ["UI-06", "Alat", "Daftar/filter/CRUD/import/export master alat", "Role sesuai menu", "KF-07 sampai KF-10"],
        ["UI-07", "Template Peminjaman", "Kelola kumpulan item intra atau antar area", "PIC, Admin, Super Admin", "KF-11"],
        ["UI-08", "Transaksi Intra Area", "Tab pengiriman dan pengembalian", "PIC, Admin, Super Admin", "KF-17 sampai KF-21"],
        ["UI-09", "Transaksi Antar Area", "Tab pengiriman/pengembalian lintas area", "PIC, Admin, Super Admin", "KF-17 sampai KF-21"],
        ["UI-10", "Mutasi Alat", "Riwayat dan aksi penerimaan/pengembalian terkait", "Role sesuai menu", "KF-18, KF-19, KF-22"],
        ["UI-11", "Riwayat Peminjaman", "Cari/filter/detail/export dan dokumen", "Semua role", "KF-16"],
        ["UI-12", "Laporan Kerusakan dan Kehilangan", "Daftar, tambah, foto, status, export, template", "Role laporan", "KF-23 sampai KF-26"],
        ["UI-13", "Pengguna", "Daftar, pencarian LDAP, role, area, CRUD", "SP/Mgr baca; Admin/Super kelola", "KF-03"],
        ["UI-14", "Area", "Daftar, cari, tambah, ubah, hapus", "Super Admin", "KF-04"],
        ["UI-15", "Activity Log dan Log Alat", "Filter, detail perubahan, dan export", "Role audit", "KF-30, KF-31"],
    ]
    add_table(doc, ["ID", "Nama Layar", "Tujuan", "Aktor", "Kebutuhan Terkait"], ui_rows, [1.2, 3.5, 7.3, 4.0, 2.6], 7.0)
    add_heading(doc, "10.2 Pola Interaksi Utama", 2)
    add_table(doc, ["Pola", "Penerapan", "Ketentuan"], [
        ["Pencarian dan filter", "Alat, peminjaman, laporan, pengguna, area, log", "Query server mengikuti area dan role; filter aktif ikut pada export jika didukung."],
        ["Keranjang dan checkout", "Peminjaman intra area", "Jumlah divalidasi ulang melalui endpoint availability sebelum penyimpanan."],
        ["Modal detail/review", "Peminjaman dan pengiriman", "Menampilkan item, status, jumlah, catatan, reviewer, dokumen, dan tindakan yang masih berlaku."],
        ["Unggah bukti", "Pengiriman, pengembalian, kerusakan, kehilangan", "Tipe, ukuran, dimensi, dan jumlah file divalidasi server."],
        ["Badge dan mailbox", "Sidebar/header", "Hanya menampilkan pekerjaan yang relevan dengan role dan area aktif."],
        ["Download", "CSV, XLSX, DOCX template", "Nama file dan akses ditentukan server; dokumen privat melalui route terautorisasi."],
    ], [4.0, 5.5, 8.5], 8.0)
    add_heading(doc, "10.3 Catatan Interaksi dan Validasi", 2)
    for number, item in enumerate([
        "Field wajib ditandai dan selalu divalidasi di server.",
        "Tombol atau menu yang tidak sesuai role harus disembunyikan, tetapi server tetap menjadi kontrol utama.",
        "Pesan validasi menjelaskan field atau item yang perlu diperbaiki tanpa menampilkan stack trace atau credential.",
        "Daftar yang panjang memakai pagination dan pencarian; state filter dipertahankan selama pengguna bekerja pada halaman.",
        "Aksi hapus, review, kirim, terima, kembali, selesai, dan perubahan status memerlukan konfirmasi yang jelas.",
        "Jumlah diminta, disetujui, telah kembali, sisa, total, dipinjam, dan tersedia harus selalu diberi label agar tidak tertukar.",
        "Status ditampilkan dengan istilah yang sama pada daftar, detail, email, export, dan Surat Jalan.",
        "Antarmuka tidak boleh menawarkan reset password atau registrasi mandiri sebelum endpoint dan proses bisnisnya disetujui.",
    ], 1):
        add_number(doc, item, number)

    # Bab 11
    add_heading(doc, "11 PEMANGKU KEPENTINGAN RACI ASUMSI BATASAN DAN KETERGANTUNGAN", 1, new_page=True)
    add_heading(doc, "11.1 Daftar Pemangku Kepentingan", 2)
    stakeholders = [
        ["Pemilik Proyek", "Sponsor, keputusan prioritas, biaya, risiko, dan penerimaan", "Tujuan, kewenangan, dan keberlanjutan layanan jelas", "Tinggi"],
        ["Pemilik Proses atau Mgr Tool", "Menetapkan aturan dan menyetujui kebutuhan", "Alur sesuai operasi alat", "Tinggi"],
        ["SP Tool", "Mereview peminjaman dan memantau laporan", "Data permintaan dan stok akurat", "Tinggi"],
        ["PIC Tool", "Mengelola alat, template, pengiriman, penerimaan, pengembalian", "Transaksi cepat dan bukti lengkap", "Tinggi"],
        ["User atau Peminjam", "Mengajukan, menerima, mengembalikan, dan memantau", "Status dan kewajiban mudah dipahami", "Sedang"],
        ["Admin dan Super Admin", "Kelola user, data, transaksi, area, dan audit", "Kontrol lengkap tetapi tetap terbatas area/role", "Tinggi"],
        ["Tim AD, Email, Jaringan", "Menyediakan LDAP, SMTP, DNS, dan sertifikat", "Integrasi stabil dan terpantau", "Tinggi"],
        ["DBA dan Administrator Server", "Oracle, storage, worker, scheduler, backup, monitoring", "Ketersediaan dan pemulihan", "Tinggi"],
        ["Tim Pengembang, QA, Keamanan", "Analisis, perubahan, test, review, release", "Perubahan terkendali tanpa regresi", "Tinggi"],
        ["Audit atau Manajemen", "Menggunakan dashboard, laporan, dan log", "Data lengkap dan dapat ditelusuri", "Sedang"],
    ]
    add_table(doc, ["Pemangku Kepentingan", "Peran atau Kontribusi", "Kebutuhan atau Tujuan", "Pengaruh"], stakeholders, [4.2, 7.0, 5.5, 1.8], 7.4)
    add_heading(doc, "11.2 Matriks RACI", 2)
    raci = [
        ["Definisi kebutuhan dan prioritas", "A", "R", "C", "C", "I", "C"],
        ["Validasi master alat dan area", "I", "A", "C", "R", "C", "C"],
        ["Pengembangan dan automated test", "I", "C", "I", "C", "I", "R/A"],
        ["UAT proses peminjaman", "I", "A", "R", "R", "R", "C"],
        ["Security review dan akses", "I", "C", "I", "C", "I", "R/A"],
        ["Kesiapan Oracle, LDAP, SMTP, storage", "I", "C", "I", "C", "I", "R/A"],
        ["Release, backup, dan rollback", "A", "C", "I", "C", "I", "R"],
        ["Operasi harian dan rekonsiliasi", "I", "A", "R", "R", "C", "C"],
        ["Approval BRD dan perubahan lingkup", "A", "R", "C", "C", "C", "C"],
    ]
    add_table(doc, ["Aktivitas", "Pemilik Proyek", "Pemilik Proses", "SP Tool", "PIC Tool", "User", "Tim TI"], raci, [5.4, 2.2, 2.3, 2.0, 2.0, 1.7, 2.4], 6.8)
    add_note(doc, "R = Responsible, A = Accountable, C = Consulted, I = Informed. Nama pemegang peran ditetapkan pada persetujuan akhir.")
    add_heading(doc, "11.3 Asumsi Batasan dan Ketergantungan", 2)
    assumptions = [
        ["Asumsi", "LDAP menyediakan username, nama, email, dan atribut yang dibutuhkan.", "Login, sinkronisasi, atau email tidak lengkap.", "Tim AD"],
        ["Asumsi", "Setiap pengguna yang boleh login telah dibuat pada master user dengan role dan area yang benar.", "Akses ditolak atau data salah cakupan.", "Admin"],
        ["Asumsi", "Data alat awal dan total aset telah direkonsiliasi oleh pemilik data.", "Availability dan laporan tidak akurat.", "Pemilik data"],
        ["Asumsi", "SMTP menerima email aplikasi dan alamat penerima valid.", "Notifikasi email gagal.", "Tim Email"],
        ["Batasan", "Kredensial tetap dikelola LDAP; aplikasi tidak mengubah password.", "Reset password mengikuti proses organisasi.", "Pemilik aplikasi"],
        ["Batasan", "Role Guest tetap merupakan akun terautentikasi dengan hak baca terbatas.", "Istilah guest dapat disalahartikan sebagai publik.", "Pemilik proses"],
        ["Batasan", "Anggaran, SLA, RPO, RTO, retensi, kapasitas, dan tanggal go live belum tersedia.", "BRD belum dapat ditandatangani final.", "Pemilik proyek"],
        ["Ketergantungan", "Oracle dan yajra/laravel-oci8 tersedia dengan konfigurasi yang benar.", "Aplikasi dan database queue tidak berjalan.", "DBA"],
        ["Ketergantungan", "LDAP/AD, DNS, CA trust, dan konektivitas internal tersedia.", "Login dan sinkronisasi gagal.", "Tim AD/Jaringan"],
        ["Ketergantungan", "Storage template Surat Jalan dan laporan tersedia serta dapat ditulis.", "Dokumen gagal dibuat atau diunduh.", "Administrator server"],
        ["Ketergantungan", "Queue worker dan scheduler dikelola sebagai service.", "Import, email, sinkronisasi, dan reminder terlambat.", "Administrator server"],
        ["Ketergantungan", "Lingkungan QA dan proses release/rollback tersedia.", "Perubahan berisiko mengganggu produksi.", "DevOps/Pemilik proyek"],
    ]
    add_table(doc, ["Jenis", "Uraian", "Dampak jika Tidak Terpenuhi", "Pemilik"], assumptions, [2.3, 8.0, 5.7, 2.0], 7.1)

    # Bab 12
    add_heading(doc, "12 UJI PENERIMAAN PENGGUNA KRITERIA PENERIMAAN DAN KETERTELUSURAN", 1, new_page=True)
    add_heading(doc, "12.1 Skenario Uji Penerimaan Pengguna", 2)
    uat = [
        ["UAT-01", "Login akun diizinkan", "Akun lokal aktif dan LDAP valid", "Login username/password", "Dashboard tampil sesuai role", "TBD UAT"],
        ["UAT-02", "Tolak akun tidak diizinkan", "LDAP valid tetapi tidak ada di master user", "Login", "Akses ditolak tanpa membuka data", "TBD UAT"],
        ["UAT-03", "Akses lintas role dan area", "Akun tiap role pada dua area", "Buka URL dan panggil API terlarang", "Server mengembalikan forbidden atau redirect aman", "TBD UAT"],
        ["UAT-04", "CRUD dan kode alat", "PIC/Admin login", "Tambah, ubah, cari, dan coba kode duplikat", "Data benar; duplikat dalam area ditolak", "TBD UAT"],
        ["UAT-05", "Import alat berhasil", "File CSV/XLSX valid", "Upload dan pantau status", "Completed; jumlah created/updated benar; log tersedia", "TBD UAT"],
        ["UAT-06", "Import alat gagal aman", "File salah format atau area", "Upload", "Failed gagal dengan pesan baris; data tidak setengah terproses", "TBD UAT"],
        ["UAT-07", "Peminjaman intra area", "Stok tersedia", "Pilih item, tanggal, pekerjaan, ajukan", "Perlu Disetujui dan stok terpesan sesuai aturan", "TBD UAT"],
        ["UAT-08", "Cegah overbooking", "Dua permintaan bersamaan", "Ajukan melebihi sisa", "Salah satu ditolak; availability tidak negatif", "TBD UAT"],
        ["UAT-09", "Review item parsial", "Pengajuan pending", "Setujui sebagian dan tolak item lain", "Jumlah dan alasan tersimpan; status benar", "TBD UAT"],
        ["UAT-10", "Peminjaman antar area", "Dua area dan stok sumber tersedia", "Ajukan, review peminjam, review sumber", "Dua reviewer/area tercatat dan jumlah tidak meningkat", "TBD UAT"],
        ["UAT-11", "Pengiriman dengan bukti", "Status Disetujui", "Unggah foto, pengirim, resi lalu kirim", "Dikirim; foto dan Surat Jalan XLSX dapat diunduh", "TBD UAT"],
        ["UAT-12", "Validasi foto", "Status Disetujui", "Unggah tipe/ukuran/jumlah di luar batas", "Ditolak dengan pesan validasi", "TBD UAT"],
        ["UAT-13", "Penerimaan intra area", "Dikirim dan user pemilik", "Klik terima", "Diterima; user lain tidak dapat menerima", "TBD UAT"],
        ["UAT-14", "Penerimaan antar area", "Dikirim lintas area", "PIC area peminjam menerima", "Diterima; stok sementara terbentuk", "TBD UAT"],
        ["UAT-15", "Pengembalian parsial", "Ada sisa item", "Kembalikan sebagian dengan foto", "Dikembalikan Partials; sisa dan Surat Jalan benar", "TBD UAT"],
        ["UAT-16", "Pengembalian seluruhnya", "Ada sisa item", "Kembalikan semua lalu selesaikan", "Dikembalikan Semuanya lalu Selesai", "TBD UAT"],
        ["UAT-17", "Laporan saat kembali", "Item kembali", "Catat rusak/hilang beserta foto", "Laporan terbentuk dan dapat diproses", "TBD UAT"],
        ["UAT-18", "Status laporan dan stok", "Laporan Dilaporkan", "Setujui, ubah ke selesai/ditolak", "Stok berkurang/terpulih tanpa duplikasi", "TBD UAT"],
        ["UAT-19", "Reminder jatuh tempo", "Status Diterima/Partials", "Jalankan dry-run H-1, H, H+6, H+12", "Target penerima dan deduplikasi harian benar", "TBD UAT"],
        ["UAT-20", "Badge dan mailbox", "Data tindakan pada beberapa area", "Login tiap role", "Hitungan hanya memuat item yang relevan", "TBD UAT"],
        ["UAT-21", "Export", "Filter aktif", "Export alat, peminjaman, laporan, dan log", "Isi sesuai filter/area dan encoding dapat dibaca", "TBD UAT"],
        ["UAT-22", "Audit trail", "Transaksi contoh lengkap", "Telusuri Activity Log dan Log Alat", "Actor, perubahan, area, waktu, dan subject konsisten", "TBD UAT"],
        ["UAT-23", "Restart operasional", "Worker dan scheduler dikonfigurasi", "Restart server dan proses", "Web, worker, scheduler, dan job kembali aktif", "TBD UAT"],
        ["UAT-24", "Backup dan restore", "Backup produksi tiruan tersedia", "Pulihkan ke lingkungan terpisah", "Data, dokumen, foto, dan relasi dapat dibaca", "TBD UAT"],
    ]
    add_table(doc, ["ID", "Skenario", "Prasyarat", "Langkah Ringkas", "Hasil yang Diharapkan", "Status"], uat, [1.3, 3.2, 3.6, 4.0, 5.4, 2.0], 6.5)
    add_heading(doc, "12.2 Kriteria Penerimaan", 2)
    acceptance = [
        ["KP-01", "Seluruh kebutuhan Wajib yang disetujui lulus UAT.", "Berita acara UAT dan bukti uji", "Pemilik Proses"],
        ["KP-02", "Tidak ada defect kritis atau tinggi terbuka pada workflow produksi.", "Defect register dan keputusan risiko", "QA dan Keamanan"],
        ["KP-03", "Test otomatis lulus pada build rilis; baseline saat penyusunan adalah 23 test dan 1.933 assertion lulus.", "Log pipeline/test rilis", "Tim Pengembang"],
        ["KP-04", "Hak akses seluruh role dan area telah diuji pada UI dan API.", "Matriks akses dan bukti negatif", "QA dan Pemilik Akses"],
        ["KP-05", "Stok, status, dokumen, foto, dan log konsisten pada sampel intra/antar area.", "Sampling database dan berkas", "Pemilik Data"],
        ["KP-06", "Oracle, LDAP, SMTP, queue, scheduler, storage, dan monitoring lulus integration check.", "Checklist integrasi dan operasional", "Tim Infrastruktur"],
        ["KP-07", "Backup dan restore berhasil; RPO/RTO disetujui.", "Laporan drill restore", "Pemilik Layanan"],
        ["KP-08", "HTTPS/DNS, retensi, dukungan, eskalasi, dan runbook tersedia.", "Checklist go live", "Pemilik Proyek"],
        ["KP-09", "BRD dan perubahan lingkup telah disetujui pihak berwenang.", "Tanda tangan atau approval resmi", "Pemilik Proyek"],
    ]
    add_table(doc, ["ID", "Kriteria Penerimaan", "Bukti atau Metode Verifikasi", "Penanggung Jawab"], acceptance, [1.4, 8.5, 5.6, 3.0], 7.5)
    add_heading(doc, "12.3 Matriks Ketertelusuran Kebutuhan", 2)
    trace = [
        ["KB-01", "KF-04 sampai KF-10", "Master alat dan dashboard", "UAT-04 sampai UAT-06, UAT-21", "Source tersedia; UAT formal TBD"],
        ["KB-02", "KF-11, KF-12, KF-15 sampai KF-21", "Siklus intra area", "UAT-07 sampai UAT-09, UAT-11, UAT-13, UAT-15, UAT-16", "Source tersedia; UAT formal TBD"],
        ["KB-03", "KF-11, KF-13 sampai KF-19, KF-21", "Siklus antar area", "UAT-10, UAT-14 sampai UAT-16", "Source tersedia; UAT formal TBD"],
        ["KB-04", "KF-01 sampai KF-04, KF-32", "Autentikasi dan akses", "UAT-01 sampai UAT-03", "Source tersedia; LDAP/UAT TBD"],
        ["KB-05", "KF-17 sampai KF-19, KF-22, KF-32", "Bukti dan mutasi", "UAT-11 sampai UAT-16", "Source tersedia; UAT formal TBD"],
        ["KB-06", "KF-23 sampai KF-26", "Laporan dan stok", "UAT-17, UAT-18", "Source tersedia; UAT formal TBD"],
        ["KB-07", "KF-05, KF-06, KF-27 sampai KF-29", "Dashboard/notifikasi", "UAT-19, UAT-20", "Source tersedia; SMTP/scheduler TBD"],
        ["KB-08", "KF-09, KF-10, KF-16, KF-26, KF-30, KF-31", "Import/export", "UAT-05, UAT-06, UAT-21", "Source tersedia; UAT formal TBD"],
        ["KB-09", "KF-16, KF-22, KF-30, KF-31", "Audit dan riwayat", "UAT-22", "Source tersedia; sampling TBD"],
        ["KB-10", "NFR operabilitas dan pemulihan", "Kesiapan produksi", "UAT-23, UAT-24", "Keputusan operasional terbuka"],
    ]
    add_table(doc, ["Kebutuhan Bisnis", "Kebutuhan Fungsional", "Desain atau Proses", "UAT", "Status"], trace, [2.6, 4.0, 3.7, 5.2, 4.5], 7.1)

    # Bab 13
    add_heading(doc, "13 JADWAL IMPLEMENTASI DAN KESIAPAN IMPLEMENTASI", 1, new_page=True)
    add_heading(doc, "13.1 Tahapan Implementasi", 2)
    schedule_rows = [
        ["1 Review Baseline", "Konfirmasi fitur pada source, test, dan manual", "Tim Pengembang", "Selesai untuk baseline 10 Sep 2026", "Daftar fitur dan hasil test"],
        ["2 Validasi BRD", "Review kebutuhan, role, area, status, owner, lingkup, dan prioritas", "Pemilik Proses", "TBD", "BRD disetujui"],
        ["3 Persiapan Data", "Cleansing area, user allow list, master alat, kode, dan stok awal", "Pemilik Data", "TBD", "Data migrasi tervalidasi"],
        ["4 QA dan Keamanan", "Regression, access matrix, concurrency, file upload, dependency review", "QA/Keamanan", "TBD", "Laporan test dan defect"],
        ["5 UAT", "Jalankan UAT-01 sampai UAT-24 dan tutup temuan", "Key User", "TBD", "Berita acara UAT"],
        ["6 Kesiapan Infrastruktur", "Oracle, LDAP, SMTP, HTTPS, storage, queue, scheduler, backup, monitoring", "Tim Infrastruktur", "TBD", "Checklist go live"],
        ["7 Release Produksi", "Backup, migration, deploy, smoke test, dan keputusan go/no-go", "DevOps/Pemilik Proyek", "TBD", "Rilis terverifikasi"],
        ["8 Stabilisasi", "Pantau error, job, email, kapasitas, data, dan feedback", "Tim Operasi", "TBD", "Laporan stabilisasi"],
    ]
    add_table(doc, ["Tahap", "Kegiatan Utama", "Penanggung Jawab", "Waktu atau Status", "Keluaran"], schedule_rows, [3.0, 6.2, 3.7, 3.1, 4.0], 7.4)
    add_heading(doc, "13.2 Kesiapan Implementasi", 2)
    readiness = [
        ["Fungsi", "Route, controller, model, service, UI, migration, mail, dan command tersedia", "Sebagian siap", "Baseline test lokal: 23 test dan 1.933 assertion lulus pada 10 Sep 2026"],
        ["Proses", "Role, status, review, pengiriman, kembali, laporan, dan audit terdokumentasi", "Perlu validasi", "Owner proses dan SOP final TBD"],
        ["Pengguna", "Tujuh role terdefinisi", "Perlu UAT/pelatihan", "Daftar user, area, dan materi training TBD"],
        ["Data", "Migration dan import alat tersedia", "Perlu verifikasi", "Cleansing dan rekonsiliasi data awal belum dibuktikan"],
        ["Teknologi", "Laravel/Vue, Oracle, LDAP, mail, storage, queue, scheduler", "Perlu gate produksi", "Konfigurasi environment, kapasitas, dan konektivitas perlu diuji"],
        ["Keamanan", "Session, CSRF, RBAC, scoping, file validation tersedia", "Perlu review formal", "HTTPS, pentest, retensi, dan access review TBD"],
        ["Operasi", "Scheduler LDAP/reminder dan contoh Supervisor worker tersedia", "Perlu implementasi", "Monitoring, backup, restore, on-call, dan eskalasi TBD"],
        ["Dokumentasi", "Manual Book dan BRD draft tersedia", "Draft", "Runbook, SOP, dan sign-off perlu diselesaikan"],
    ]
    add_table(doc, ["Area Kesiapan", "Kriteria", "Status", "Bukti atau Catatan"], readiness, [3.0, 7.3, 2.8, 6.7], 7.4)
    add_heading(doc, "13.3 Rencana Implementasi Produksi dan Pasca Implementasi", 2)
    for number, item in enumerate([
        "Bekukan kebutuhan perubahan dan ambil backup terverifikasi sebelum migration atau deployment.",
        "Pastikan konfigurasi Production terpisah dari QA dan credential tidak tersimpan di source.",
        "Jalankan migration review, automated test, security gate, dan UAT pada build yang sama dengan kandidat rilis.",
        "Aktifkan worker database queue dengan Supervisor atau service manager serta Laravel scheduler.",
        "Lakukan smoke test login LDAP, dashboard, master alat, peminjaman, review, Surat Jalan, email, laporan, export, dan log.",
        "Pantau Oracle, error log, failed jobs, scheduler, email, storage, dan kapasitas selama masa stabilisasi.",
        "Gunakan rollback plan jika smoke test gagal dan rekonsiliasi setiap perubahan data yang sudah terjadi.",
        "Laksanakan post implementation review dan masukkan perubahan yang disetujui ke change log BRD.",
    ], 1):
        add_number(doc, item, number)

    # Bab 14
    add_heading(doc, "14 PERSETUJUAN AKHIR DAN PENGENDALIAN PERUBAHAN", 1, new_page=True)
    add_heading(doc, "14.1 Pernyataan Persetujuan", 2)
    add_body(doc, "Dengan menandatangani bagian ini, para pihak menyatakan bahwa kebutuhan bisnis pada dokumen telah ditinjau dan disepakati sebagai acuan pelaksanaan IPTOOLS LITE. Perubahan setelah persetujuan mengikuti mekanisme pengendalian perubahan.")
    add_table(doc, ["Peran", "Nama atau Jabatan", "Keputusan", "Tanggal", "Tanda Tangan"], [
        ["Pemilik Proyek atau Penanggung Jawab Utama", "[Nama]", "[Setuju atau Tidak]", "", ""],
        ["Pemilik Proses Bisnis", "[Nama]", "[Setuju atau Tidak]", "", ""],
        ["Fungsi TI atau Solusi", "[Nama]", "[Setuju atau Tidak]", "", ""],
    ], [4.2, 5.7, 3.0, 2.4, 3.0], 8.1)
    add_heading(doc, "14.2 Pengendalian Perubahan Kebutuhan", 2)
    add_table(doc, ["ID Perubahan", "Tanggal", "Pemohon", "Uraian Perubahan", "Dampak", "Keputusan"], [
        ["CR-001", "[Tanggal]", "[Pemohon]", "[Perubahan]", "[Lingkup, waktu, biaya, risiko]", "[Setuju, tolak, tunda]"],
        ["CR-002", "", "", "", "", ""],
    ], [2.2, 2.4, 2.7, 5.8, 4.1, 2.8], 7.5)
    add_heading(doc, "14.3 Keputusan yang Masih Terbuka", 2)
    for number, item in enumerate([
        "Nama pemilik proyek, manajer proyek, pemilik proses, dan approver BRD.",
        "Tanggal go live, anggaran, kapasitas, serta sumber daya operasi dan dukungan.",
        "SLA, RPO, RTO, jadwal backup, lokasi salinan, retensi database, foto, dokumen, dan log.",
        "DNS, sertifikat HTTPS, segmentasi jaringan, dan daftar browser yang didukung.",
        "Konfigurasi queue production, Supervisor/service manager, scheduler, monitoring, serta eskalasi.",
        "Daftar area, mapping role, owner data, data awal, dan hasil rekonsiliasi stok.",
        "Hasil security review, concurrency test, UAT formal, restore drill, dan persetujuan risiko residual.",
        "Keputusan apakah role Guest tetap dipakai untuk akun baca terbatas atau diganti namanya agar tidak dianggap akses publik.",
    ], 1):
        add_number(doc, item, number)

    add_heading(doc, "Lampiran A Daftar Istilah", 2)
    glossary = [
        ["Activity Log", "Catatan aktivitas aplikasi yang menyimpan actor, route, area, perubahan, dan metadata request."],
        ["Alat", "Jenis/perangkat yang dikelola per area, dengan kode, klasifikasi, lokasi, dan total aset."],
        ["AreaAlatStock", "Stok sementara di area peminjam yang berasal dari transaksi antar area."],
        ["LDAP", "Lightweight Directory Access Protocol yang dipakai untuk verifikasi kredensial dan sinkronisasi user."],
        ["PIC Tool", "Petugas pengelola alat, pengiriman, penerimaan, dan pengembalian pada area."],
        ["SP Tool", "Role reviewer peminjaman dan pengelola proses tertentu sesuai area."],
        ["Mgr Tool", "Role manajerial untuk review, monitoring, dan audit sesuai akses."],
        ["Availability", "Jumlah yang masih dapat dipinjam setelah mengurangi kuantitas pada status pemegang stok."],
        ["Intra Area", "Peminjaman dalam area yang sama."],
        ["Antar Area", "Peminjaman dengan area peminjam berbeda dari area sumber."],
        ["Surat Jalan", "Dokumen XLSX yang dihasilkan untuk pengiriman atau pengembalian dan dapat memuat lampiran foto."],
        ["RACI", "Responsible, Accountable, Consulted, dan Informed."],
        ["RPO dan RTO", "Recovery Point Objective dan Recovery Time Objective."],
        ["UAT", "User Acceptance Test atau uji penerimaan pengguna."],
    ]
    add_table(doc, ["Istilah", "Definisi"], glossary, [4.0, 14.0], 8.2)
    add_heading(doc, "Lampiran B Sumber Penyusunan", 2)
    sources = [
        "BRD InfraSight.docx sebagai referensi struktur dan tata letak.",
        "Source code IPTOOLS LITE pada repository laravel-inertia: route, controller, middleware, model, service, job, mail, command, migration, configuration, Vue page/component, dan test.",
        "Manual_Book_IPTOOLS_LITE.docx.08b626c3 sebagai sumber istilah dan gambaran penggunaan; fakta yang bertentangan dengan route/controller tidak dipakai.",
        "Hasil test lokal pada 10 September 2026: 23 test lulus dengan 1.933 assertion.",
        "README produksi mengenai database queue dan rekomendasi Supervisor worker.",
    ]
    for item in sources:
        add_bullet(doc, item)

    doc.core_properties.title = "Business Requirements Document IPTOOLS LITE"
    doc.core_properties.subject = "Kebutuhan bisnis dan fungsional sistem pengelolaan alat dan peminjaman antar area"
    doc.core_properties.author = "Tim IPTOOLS LITE"
    doc.core_properties.keywords = "BRD, IPTOOLS LITE, peminjaman alat, inventaris, Laravel, Oracle, LDAP"

    doc.save(OUTPUT)
    patch_header_text(OUTPUT)
    print(str(OUTPUT))


if __name__ == "__main__":
    build()
