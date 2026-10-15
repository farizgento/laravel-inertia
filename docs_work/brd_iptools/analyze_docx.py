from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def length_pt(value):
    return None if value is None else round(value.pt, 2)


def length_cm(value):
    return None if value is None else round(value.cm, 3)


def color_value(color):
    if color is None or color.rgb is None:
        return None
    return str(color.rgb)


def paragraph_info(p):
    fmt = p.paragraph_format
    runs = []
    for r in p.runs:
        if not r.text and r._element.xpath('.//w:drawing|.//w:pict'):
            text = '[IMAGE]'
        else:
            text = r.text
        runs.append({
            'text': text,
            'bold': r.bold,
            'italic': r.italic,
            'underline': bool(r.underline) if r.underline is not None else None,
            'font': r.font.name,
            'size_pt': length_pt(r.font.size),
            'color': color_value(r.font.color),
        })
    return {
        'text': p.text,
        'style': p.style.name if p.style else None,
        'alignment': str(p.alignment),
        'space_before_pt': length_pt(fmt.space_before),
        'space_after_pt': length_pt(fmt.space_after),
        'line_spacing': str(fmt.line_spacing),
        'left_indent_cm': length_cm(fmt.left_indent),
        'first_line_indent_cm': length_cm(fmt.first_line_indent),
        'keep_with_next': fmt.keep_with_next,
        'page_break_before': fmt.page_break_before,
        'runs': runs,
    }


def table_info(table):
    grid = []
    for row in table.rows:
        grid.append([cell.text for cell in row.cells])
    widths = []
    if table.rows:
        for cell in table.rows[0].cells:
            widths.append(length_cm(cell.width))
    try:
        col_count = len(table.columns)
    except Exception:
        col_count = max((len(row.cells) for row in table.rows), default=0)
    return {
        'style': table.style.name if table.style else None,
        'rows': len(table.rows),
        'cols': col_count,
        'first_row_widths_cm': widths,
        'grid': grid,
    }


def main(path_s: str):
    path = Path(path_s)
    doc = Document(path)
    out = {
        'path': str(path.resolve()),
        'size': path.stat().st_size,
        'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
        'core_properties': {
            'title': doc.core_properties.title,
            'subject': doc.core_properties.subject,
            'author': doc.core_properties.author,
            'created': str(doc.core_properties.created),
            'modified': str(doc.core_properties.modified),
        },
        'sections': [],
        'paragraphs': [],
        'tables': [],
        'headers': [],
        'footers': [],
        'style_counts': Counter(),
        'inline_shapes': len(doc.inline_shapes),
        'style_definitions': {},
    }
    for style_name in [
        'Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2',
        'List Paragraph', 'Header', 'Footer', 'TOC 1', 'TOC 2',
    ]:
        if style_name not in [s.name for s in doc.styles]:
            continue
        style = doc.styles[style_name]
        fmt = style.paragraph_format
        out['style_definitions'][style_name] = {
            'font': style.font.name,
            'size_pt': length_pt(style.font.size),
            'bold': style.font.bold,
            'italic': style.font.italic,
            'color': color_value(style.font.color),
            'alignment': str(fmt.alignment),
            'space_before_pt': length_pt(fmt.space_before),
            'space_after_pt': length_pt(fmt.space_after),
            'line_spacing': str(fmt.line_spacing),
            'keep_with_next': fmt.keep_with_next,
        }
    for sec in doc.sections:
        out['sections'].append({
            'page_width_cm': length_cm(sec.page_width),
            'page_height_cm': length_cm(sec.page_height),
            'orientation': str(sec.orientation),
            'top_margin_cm': length_cm(sec.top_margin),
            'bottom_margin_cm': length_cm(sec.bottom_margin),
            'left_margin_cm': length_cm(sec.left_margin),
            'right_margin_cm': length_cm(sec.right_margin),
            'header_distance_cm': length_cm(sec.header_distance),
            'footer_distance_cm': length_cm(sec.footer_distance),
            'start_type': str(sec.start_type),
            'different_first_page_header_footer': sec.different_first_page_header_footer,
        })
        out['headers'].append([paragraph_info(p) for p in sec.header.paragraphs])
        out['footers'].append([paragraph_info(p) for p in sec.footer.paragraphs])
    for p in doc.paragraphs:
        info = paragraph_info(p)
        out['paragraphs'].append(info)
        out['style_counts'][info['style']] += 1
    for t in doc.tables:
        out['tables'].append(table_info(t))
    out['style_counts'] = dict(out['style_counts'])
    with zipfile.ZipFile(path) as z:
        out['package_parts'] = [
            {
                'name': i.filename,
                'size': i.file_size,
                'sha256': hashlib.sha256(z.read(i.filename)).hexdigest(),
            }
            for i in z.infolist()
        ]
        xml = z.read('word/document.xml')
        out['document_xml_text_node_count'] = xml.count(b'<w:t')
        out['has_fields'] = b'w:fldChar' in xml or b'w:instrText' in xml
        out['has_textboxes'] = b'w:txbxContent' in xml
        out['has_content_controls'] = b'<w:sdt' in xml
    print(json.dumps(out, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
    main(sys.argv[1])
